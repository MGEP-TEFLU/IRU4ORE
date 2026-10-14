# -*- coding: utf-8 -*-
"""
Created on Tue Oct  6 13:32:12 2026
Module name: [Resource_indicators_calculator]
Author: [Ander Martinez, Markel Penalba]
Affiliation: [Mondragon University]
Contact: [amartinezpe@mondragon.edu]
"""



'--> Dependencies'

# Python libraries
import pandas as pd
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
# Personal functions
from resource_fcn_IRU4ORE import *#OMAE2020,env_contour_plot,wawi_energy_flux,scatter_diagram,rank_cor


def get_met_data(init_y,end_y,locations,rooth_dir):
    
    met_data = {}                                                               # dic file containing metocean data from the selected sites
    
    for i in range(len(locations)):
        met_data[locations[i]] = pd.read_csv(rooth_dir+'/Metocean_data/'+locations[i]+"_calibdata_3m10m.csv.gz", compression="gzip")
        met_data[locations[i]] = met_data[locations[i]].set_index(['Datetime']) 
        met_data[locations[i]] = met_data[locations[i]][str(init_y[i]):str(end_y[i])]
        
        
    return met_data



def res_in_calc(met_data,locations,init_y,end_y,rooth_dir):
    
    '--> Dic file where resource indicators are saved'
    res_param = {}
    sources = ["ERA5", "Buoy", "CERA5"]
    variables = ["Wave", "WaveWind"]
    extreme_types = ["EDP", "Contour"]
    
    res_param = {
        "Resource": {
            "rank_correlation":{},
            "extremes": {
                var: {
                    ext: {src: {} for src in sources}
                    for ext in extreme_types
                }
                for var in variables
            },
            "flux_energy": {},
            "Scatter_Wave": {src: {} for src in sources},
            "Scatter_WaveWind": {src: {} for src in sources}
        },
        "Resource_unc":{
            "EDP_unc":{},
            "Joint_prob_unc":{},
            "flux_unc":{},            
            
            }
    }
    
    
    
    # Specify which parameters have to be calculated
    rank_flag = True
    scatter_dig_flag = True
    energy_flux_flag = True
    env_cont_UwHs_flag = True
    env_cont_TpHs_flag = True
    plot_env_cont = False
    plot_energy_flux = False
    plot_scatter = False
    
    
    '--> Per site resource indicators calculation'    
    for i in range(len(locations)):
        
        '0.- Spearman´s Rank correlation'
        if rank_flag == True:
            res_param["Resource"]["rank_correlation"][locations[i]] = rank_cor(met_data,locations[i])
        
        '1.- IFORM environmental contour creation'
        tr = 50                                                                                       # Return years 
        ts = 1                                                                                        # Sampling interval [h]
        
        '1.1.- Hs vs Tp IFORM'
        if env_cont_TpHs_flag == True:
            res_param["Resource"]["extremes"]["Wave"]["EDP"]["Buoy"][locations[i]],res_param["Resource"]["extremes"]["Wave"]["Contour"]["Buoy"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Obs","Tp_Obs"]],tr,ts,locations[i],'Buoy',plot_env_cont)   
            res_param["Resource"]["extremes"]["Wave"]["EDP"]["ERA5"][locations[i]],res_param["Resource"]["extremes"]["Wave"]["Contour"]["ERA5"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Sim","Tp_Sim"]],tr,ts,locations[i],'ERA5',plot_env_cont)
            res_param["Resource"]["extremes"]["Wave"]["EDP"]["CERA5"][locations[i]],res_param["Resource"]["extremes"]["Wave"]["Contour"]["CERA5"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Calib","Tp_Calib"]],tr,ts,locations[i],'CERA5',plot_env_cont)
        
        '1.2.- Hs vs Uw IFORM'
        if env_cont_UwHs_flag == True:
            res_param["Resource"]["extremes"]["WaveWind"]["EDP"]["Buoy"][locations[i]],res_param["Resource"]["extremes"]["WaveWind"]["Contour"]["Buoy"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Obs","Uv10_Obs"]],tr,ts,locations[i],'Buoy',plot_env_cont)   
            res_param["Resource"]["extremes"]["WaveWind"]["EDP"]["ERA5"][locations[i]],res_param["Resource"]["extremes"]["WaveWind"]["Contour"]["ERA5"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Sim","Uw_Sim"]],tr,ts,locations[i],'ERA5',plot_env_cont)
            res_param["Resource"]["extremes"]["WaveWind"]["EDP"]["CERA5"][locations[i]],res_param["Resource"]["extremes"]["WaveWind"]["Contour"]["CERA5"][locations[i]]  =  OMAE2020(met_data[locations[i]][["Hs_Calib","Uv10_Calib"]],tr,ts,locations[i],'CERA5',plot_env_cont)
            
        '1.3.- Plotting IFORMs'
        if plot_env_cont == True:
            env_contour_plot(res_param,locations[i],rooth_dir)
            
            
        '2.- Wave and wind energy flux'
        if energy_flux_flag == True:
            
            res_param["Resource"]["flux_energy"][locations[i]] = wawi_energy_flux(met_data,locations[i],rooth_dir,plot_energy_flux)
            
        '3.- Wave and wind scatter diagram'
        if scatter_dig_flag  == True:
              res_param = scatter_diagram(met_data[locations[i]],res_param,init_y[i],end_y[i],locations[i],scatter_dig_flag)
              
        plt.close("all")
        
        
    return res_param
              

def res_ind_uncert_cal(res_param,locations,root_dir):
    
    '--> Per site and resource indicator uncertainty calculation'    
    for i in range(len(locations)):    
    
        '1.- Extreme design point discrepancy (derived from IFORM)'
        
        res_param["Resource_unc"]["EDP_unc"][locations[i]] = uncert_EDP(res_param,locations[i])
        
        '2.- Wave and wind flux error'
        res_param["Resource_unc"]["flux_unc"][locations[i]] = uncert_flux(res_param,locations[i])
        
        '3.- Occurrence error'
        res_param["Resource_unc"]["Joint_prob_unc"][locations[i]] = uncert_occur(res_param,locations[i],True,root_dir)
        
    return res_param
        
        
        
def save_res_ind(res_param,locations,root_dir):
    
    # ============================================================
    # SETTINGS
    # ============================================================
    path_res = root_dir+"/Results/Resource_indicators/"
    os.makedirs(path_res, exist_ok=True) 
    output_path = Path(path_res+"Resource_indicators.xlsx")
   
        
    datasets = ["Buoy", "ERA5", "CERA5"]
    rank_datasets = ["ERA5", "CERA5"]
    rank_datasets_v2 = ["Sim", "Calib"]
    
    flux_dataset_map = {"Buoy": "obs", "ERA5": "sim", "CERA5": "calib"}
    
    locations = list(res_param["Resource"]["flux_energy"].keys())
    
    extremes_dict = res_param["Resource"]["extremes"]
    wind_key = "Wind" if "Wind" in extremes_dict else "WaveWind"
    
    # ============================================================
    # AUXILIARY FUNCTIONS
    # ============================================================
    def to_scalar(value):
        arr = np.asarray(value).squeeze()
        if np.size(arr) != 1: raise ValueError(f"Expected a scalar value, but obtained shape {np.shape(arr)}")
        return float(arr)
    
    def get_df_column_value(df, column):
        return to_scalar(df[column])
    
    def write_two_level_table(sheet, row_labels, locations, subcolumns, values):
        n_sub = len(subcolumns)
    
        sheet.merge_cells("A1:A2")
        sheet["A1"] = "Parameter"
        sheet["A1"].font = Font(bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
        col = 2
    
        for location in locations:
            start_col = col
            end_col = col + n_sub - 1
            sheet.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    
            cell = sheet.cell(row=1, column=start_col)
            cell.value = location
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
            for j, dataset in enumerate(subcolumns):
                cell = sheet.cell(row=2, column=start_col + j)
                cell.value = dataset
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
            col += n_sub
    
        for i, label in enumerate(row_labels, start=3):
            sheet.cell(row=i, column=1).value = label
            sheet.cell(row=i, column=1).font = Font(bold=True)
    
            col = 2
    
            for location in locations:
                for dataset in subcolumns:
                    value = values[label][location].get(dataset, None)
                    cell = sheet.cell(row=i, column=col)
                    cell.value = value
                    if isinstance(value, (int, float, np.integer, np.floating)): cell.number_format = "0.000"
                    col += 1
    
        max_col = 1 + len(locations) * n_sub
        max_row = 2 + len(row_labels)
    
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                if cell.column != 1: cell.alignment = Alignment(horizontal="center", vertical="center")
    
        sheet.column_dimensions["A"].width = 28
    
        for col_idx in range(2, max_col + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 14
    
        sheet.row_dimensions[1].height = 22
        sheet.row_dimensions[2].height = 20
        sheet.freeze_panes = "B3"
    
    # ============================================================
    # 1. EXTREMES
    # ============================================================
    extreme_rows = ["Wave EDP - Hs [m]", "Wave EDP - Tp [s]", "Wind EDP - Hs [m]", "Wind EDP - Uw [m/s]"]
    
    extreme_values = {row: {location: {} for location in locations} for row in extreme_rows}
    
    for location in locations:
        for dataset in datasets:
            wave_edp = res_param["Resource"]["extremes"]["Wave"]["EDP"][dataset][location]
            extreme_values["Wave EDP - Hs [m]"][location][dataset] = get_df_column_value(wave_edp, "Hs_extrem")
            extreme_values["Wave EDP - Tp [s]"][location][dataset] = get_df_column_value(wave_edp, "Tp_extrem")
    
            wind_edp = res_param["Resource"]["extremes"][wind_key]["EDP"][dataset][location]
            extreme_values["Wind EDP - Hs [m]"][location][dataset] = get_df_column_value(wind_edp, "Hs_extrem")
            extreme_values["Wind EDP - Uw [m/s]"][location][dataset] = get_df_column_value(wind_edp, "Uw_extrem")
    
    # ============================================================
    # 2. FLUX ENERGY
    # ============================================================
    flux_rows = ["Wave MEFD [deg]", "Wind MEFD [deg]"]
    
    flux_values = {row: {location: {} for location in locations} for row in flux_rows}
    
    for location in locations:
        flux_df = res_param["Resource"]["flux_energy"][location]
    
        for dataset in datasets:
            suffix = flux_dataset_map[dataset]
            flux_values["Wave MEFD [deg]"][location][dataset] = get_df_column_value(flux_df, f"MEFD_wave_{suffix}")
            flux_values["Wind MEFD [deg]"][location][dataset] = get_df_column_value(flux_df, f"MEFD_wind_{suffix}")
    
    # ============================================================
    # 3. RANK CORRELATION
    # ============================================================
    rank_rows = ["Spearman rho (Hs, Tp)", "Spearman rho (Hs, Uw)"]
    
    rank_values = {row: {location: {} for location in locations} for row in rank_rows}
    
    for location in locations:
        rho_df = res_param["Resource"]["rank_correlation"][location]
    
        for i_dat, dataset in enumerate(rank_datasets):
            internal_dataset = rank_datasets_v2[i_dat]
            rank_values["Spearman rho (Hs, Tp)"][location][dataset] = to_scalar(rho_df.loc[internal_dataset, "Hs&Tp"])
            rank_values["Spearman rho (Hs, Uw)"][location][dataset] = to_scalar(rho_df.loc[internal_dataset, "Hs&Uw"])
    
    # ============================================================
    # CREATE EXCEL WORKBOOK
    # ============================================================
    wb = Workbook()
    
    sheet_extremes = wb.active
    sheet_extremes.title = "extremes"
    
    sheet_flux = wb.create_sheet("flux_energy")
    sheet_rank = wb.create_sheet("rank_correlation")
    
    # ============================================================
    # WRITE SHEETS
    # ============================================================
    write_two_level_table(sheet_extremes, extreme_rows, locations, datasets, extreme_values)
    write_two_level_table(sheet_flux, flux_rows, locations, datasets, flux_values)
    write_two_level_table(sheet_rank, rank_rows, locations, rank_datasets, rank_values)
    
    # ============================================================
    # SAVE WORKBOOK
    # ============================================================
    wb.save(output_path)
    
            
        
        
        
def save_un_res(res_param,locations,root_dir):
    
    # ============================================================
    # SETTINGS
    # ============================================================
    path_res = root_dir+"/Results/Resource_indicators/"
    os.makedirs(path_res, exist_ok=True) 
    output_path = Path(path_res+"Uncertainty_resource_indicators.xlsx")
    datasets = ["ERA5", "CERA5"]
    locations = list(res_param["Resource_unc"]["EDP_unc"].keys())
    
    # ============================================================
    # AUXILIARY FUNCTIONS
    # ============================================================
    def to_scalar(value):
        arr = np.asarray(value).squeeze()
        if np.size(arr) != 1: raise ValueError(f"Expected a scalar value, but obtained shape {np.shape(arr)}")
        return float(arr)
    
    def write_two_level_table(sheet, row_labels, locations, subcolumns, values):
        n_sub = len(subcolumns)
    
        sheet.merge_cells("A1:A2")
        sheet["A1"] = "Parameter"
        sheet["A1"].font = Font(bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
        col = 2
    
        for location in locations:
            start_col = col
            end_col = col + n_sub - 1
            sheet.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    
            cell = sheet.cell(row=1, column=start_col)
            cell.value = location
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
            for j, dataset in enumerate(subcolumns):
                cell = sheet.cell(row=2, column=start_col + j)
                cell.value = dataset
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
            col += n_sub
    
        for i, label in enumerate(row_labels, start=3):
            sheet.cell(row=i, column=1).value = label
            sheet.cell(row=i, column=1).font = Font(bold=True)
    
            col = 2
    
            for location in locations:
                for dataset in subcolumns:
                    value = values[label][location].get(dataset, None)
                    cell = sheet.cell(row=i, column=col)
                    cell.value = value
                    if isinstance(value, (int, float, np.integer, np.floating)): cell.number_format = "0.000"
                    col += 1
    
        max_col = 1 + len(locations) * n_sub
        max_row = 2 + len(row_labels)
    
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                if cell.column != 1: cell.alignment = Alignment(horizontal="center", vertical="center")
    
        sheet.column_dimensions["A"].width = 28
    
        for col_idx in range(2, max_col + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 14
    
        sheet.freeze_panes = "B3"
    
    # ============================================================
    # FUNCTION TO CONVERT SITE DATAFRAMES TO EXCEL STRUCTURE
    # ============================================================
    def extract_uncertainty_values(section_name):
        row_labels = list(res_param["Resource_unc"][section_name][locations[0]].columns)
        values = {row: {location: {} for location in locations} for row in row_labels}
    
        for location in locations:
            df = res_param["Resource_unc"][section_name][location]
    
            for dataset in datasets:
                for row in row_labels:
                    values[row][location][dataset] = to_scalar(df.loc[dataset, row])
    
        return row_labels, values
    
    # ============================================================
    # 1. EDP UNCERTAINTY
    # ============================================================
    EDP_rows, EDP_values = extract_uncertainty_values("EDP_unc")
    
    # ============================================================
    # 2. FLUX UNCERTAINTY
    # ============================================================
    flux_rows, flux_values = extract_uncertainty_values("flux_unc")
    
    # ============================================================
    # 3. JOINT PROBABILITY UNCERTAINTY
    # ============================================================
    joint_rows, joint_values = extract_uncertainty_values("Joint_prob_unc")
    
    # ============================================================
    # CREATE WORKBOOK
    # ============================================================
    wb = Workbook()
    
    sheet_EDP = wb.active
    sheet_EDP.title = "EDP_unc"
    
    sheet_flux = wb.create_sheet("flux_unc")
    sheet_joint = wb.create_sheet("Joint_prob_unc")
    
    # ============================================================
    # WRITE SHEETS
    # ============================================================
    write_two_level_table(sheet_EDP, EDP_rows, locations, datasets, EDP_values)
    write_two_level_table(sheet_flux, flux_rows, locations, datasets, flux_values)
    write_two_level_table(sheet_joint, joint_rows, locations, datasets, joint_values)
    
    # ============================================================
    # SAVE
    # ============================================================
    wb.save(output_path)
    
           
        
              
                      
                                                          
                    
                
               